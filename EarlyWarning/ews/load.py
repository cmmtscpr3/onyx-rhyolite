"""Step 1 -- read the source files exactly as they are.

Nothing is cleaned or judged here; ``validate.check_loaded`` does that.  The
output keeps every reading with its original date and text, so
``data/steps/01_loaded.csv`` shows the data as the files hold it.

* PIHPS: 3 markets x 8 yearly workbooks, the 10 commodity **group** rows only.
* BI SPIP: ``bi_card_transactions.csv``, ``bi_emoney.csv``, ``bi_payment_system.csv``.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

from . import paths
from .inventory import SPIP_FILES

paths.collectors_importable()

from collectors.sinks.long_csv import read_csv as read_long_csv  # noqa: E402
from collectors.sinks.wide_xlsx import is_value, parse_header_date, read_grid  # noqa: E402

#: Folder name -> market part of the series id.
MARKETS: dict[str, str] = {
    "Traditional Market": "traditional",
    "Modern Market": "modern",
    "Wholesale": "wholesale",
}

#: The 10 PIHPS commodity groups, numbered 1-10, as (name in the file, series id part).
GROUPS: tuple[tuple[str, str], ...] = (
    ("Beras", "01_beras"),
    ("Daging Ayam", "02_daging_ayam"),
    ("Daging Sapi", "03_daging_sapi"),
    ("Telur Ayam", "04_telur_ayam"),
    ("Bawang Merah", "05_bawang_merah"),
    ("Bawang Putih", "06_bawang_putih"),
    ("Cabai Merah", "07_cabai_merah"),
    ("Cabai Rawit", "08_cabai_rawit"),
    ("Minyak Goreng", "09_minyak_goreng"),
    ("Gula Pasir", "10_gula_pasir"),
)

_WHITESPACE = re.compile(r"\s+")
_WORKBOOK_YEAR = re.compile(r"(\d{4})\.xlsx$")

LOADED_COLUMNS = (
    "series_id", "source_date", "raw_text", "raw_value", "unit", "frequency", "source_file", "file_year",
)


def normalise(name: object) -> str:
    """Case- and space-insensitive name, so ``'Cabai Merah Keriting '`` matches."""
    return _WHITESPACE.sub(" ", str(name or "")).strip().casefold()


GROUP_IDS: dict[str, str] = {normalise(name): part for name, part in GROUPS}


@dataclass
class PihpsFile:
    """What a PIHPS workbook's layout looked like, for the structure checks."""

    path: str
    market: str
    file_year: int
    commodity_header: str
    bad_header_cells: list[str] = field(default_factory=list)
    group_names: list[str] = field(default_factory=list)


@dataclass
class Loaded:
    frame: pd.DataFrame
    pihps_files: list[PihpsFile]
    #: series ids found in the SPIP files that are not in the inventory
    extra_ids: set[str]


def parse_number(text: object) -> float:
    """``'15,750'`` -> 15750.0; a dash or blank -> NaN; anything else non-numeric -> NaN.

    The caller keeps ``raw_text`` so check Z5 can tell "missing" from "not a number".
    """
    if not is_value(text):
        return float("nan")
    try:
        return float(str(text).replace(",", "").strip())
    except ValueError:
        return float("nan")


def _read_header_row(path: Path) -> list[object]:
    """Row 1 of a PIHPS workbook.

    ``read_grid`` skips header cells it cannot parse, so the raw header is read
    separately for checks U2 and S2.  ``reset_dimensions`` is needed because
    these files declare a stale ``A1:C1`` sheet size.
    """
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[book.sheetnames[0]]
        sheet.reset_dimensions()
        for row in sheet.iter_rows(min_row=1, max_row=1, values_only=True):
            return list(row)
        return []
    finally:
        book.close()


def load_pihps(inventory: pd.DataFrame, root: Path = paths.FOOD_PRICES) -> tuple[list[dict], list[PihpsFile]]:
    records: list[dict] = []
    files: list[PihpsFile] = []
    for folder, market in MARKETS.items():
        for path in sorted((root / folder).glob("*.xlsx")):
            match = _WORKBOOK_YEAR.search(path.name)
            file_year = int(match.group(1)) if match else -1
            header = _read_header_row(path)
            info = PihpsFile(
                path=paths.relative(path),
                market=market,
                file_year=file_year,
                commodity_header="" if len(header) < 2 or header[1] is None else str(header[1]),
            )
            for cell in header[2:]:
                if cell is not None and str(cell).strip() and parse_header_date(cell) is None:
                    info.bad_header_cells.append(str(cell))
            grid = read_grid(path)
            for row in grid.rows:
                if row.level != 1:
                    continue
                info.group_names.append(row.name)
                part = GROUP_IDS.get(normalise(row.name))
                if part is None:
                    continue  # reported by check S1
                series_id = f"pihps.{market}.{part}"
                meta = inventory.loc[series_id] if series_id in inventory.index else None
                for week in grid.weeks:
                    text = grid.cell(row, week)
                    records.append({
                        "series_id": series_id,
                        "source_date": pd.Timestamp(week),
                        "raw_text": text,
                        "raw_value": parse_number(text),
                        "unit": meta["unit"] if meta is not None else "",
                        "frequency": "weekly",
                        "source_file": info.path,
                        "file_year": file_year,
                    })
            files.append(info)
    return records, files


def load_spip(inventory: pd.DataFrame, root: Path = paths.CONSUMPTION) -> tuple[list[dict], set[str]]:
    records: list[dict] = []
    extra: set[str] = set()
    for name in SPIP_FILES.values():
        path = root / name
        _, rows = read_long_csv(path)
        for row in rows:
            series_id = row["series_id"]
            if series_id not in inventory.index:
                extra.add(series_id)
                continue
            records.append({
                "series_id": series_id,
                "source_date": pd.to_datetime(row["ref_date"], format="%Y-%m-%d", errors="coerce"),
                "raw_text": row["value"],
                "raw_value": parse_number(row["value"]),
                "unit": row.get("unit", ""),
                "frequency": inventory.loc[series_id, "frequency"],
                "source_file": paths.relative(path),
                "file_year": -1,
            })
    return records, extra


def load_all(inventory: pd.DataFrame) -> Loaded:
    pihps_records, pihps_files = load_pihps(inventory)
    spip_records, extra = load_spip(inventory)
    frame = pd.DataFrame.from_records(pihps_records + spip_records, columns=list(LOADED_COLUMNS))
    frame = frame.sort_values(["series_id", "source_date", "source_file"], kind="stable").reset_index(drop=True)
    return Loaded(frame=frame, pihps_files=pihps_files, extra_ids=extra)
