"""Reading the workbooks BI, OJK and the BPS-sourced tables are published in.

Two formats and one recurring layout.  BI still publishes BIFF8 ``.xls``;
OJK and the BI survey data-series use ``.xlsx``.  Both lay their time axis out
the same way: one header row of years, each spanning a block, and one header
row below it of period labels — twelve months, or four quarters, with the
block's own total in the column where that label is blank.

``periods`` turns those two rows into a column-to-date map, which is what every
workbook collector actually needs.  Discovering the header rows rather than
hard-coding them matters: BI reshapes these tables without notice, and a
fixed cell position fails silently by reading the wrong year.
"""

from __future__ import annotations

import datetime as dt
import io
from dataclasses import dataclass
from typing import Any, Iterator

from .dates import (
    clean_label,
    month_start,
    parse_month,
    parse_quarter,
    parse_year,
    quarter_start,
)


class WorkbookError(ValueError):
    """The workbook is not shaped the way the collector expects."""


@dataclass(frozen=True, slots=True)
class Grid:
    """One sheet, materialised as rows of values."""

    name: str
    rows: tuple[tuple[Any, ...], ...]

    @property
    def nrows(self) -> int:
        return len(self.rows)

    @property
    def ncols(self) -> int:
        return max((len(row) for row in self.rows), default=0)

    def cell(self, row: int, col: int) -> Any:
        if 0 <= row < len(self.rows) and 0 <= col < len(self.rows[row]):
            return self.rows[row][col]
        return None

    def label(self, row: int, col: int) -> str:
        return clean_label(self.cell(row, col))

    def row_label(self, row: int, max_col: int = 6) -> str:
        """The first non-empty text cell in a row — the component name."""
        for col in range(min(max_col, self.ncols)):
            text = self.label(row, col)
            if text and not _is_number(self.cell(row, col)):
                return text
        return ""

    def find_row(self, *needles: str, max_col: int = 6, start: int = 0) -> int | None:
        """First row whose leading cells contain all of ``needles``."""
        wanted = [needle.casefold() for needle in needles]
        for row in range(start, self.nrows):
            haystack = " ".join(
                self.label(row, col) for col in range(min(max_col, self.ncols))
            ).casefold()
            if all(needle in haystack for needle in wanted):
                return row
        return None


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


#: Magic bytes: OLE2 compound document (legacy .xls) and zip (.xlsx/.xlsm).
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ZIP_MAGIC = b"PK\x03\x04"


def sniff(content: bytes, suffix: str = "") -> str:
    """Workbook format from the bytes, falling back to the file extension.

    The extension lies often enough to matter: BI serves ``.xls`` URLs, and a
    publisher that switches to the modern format without renaming the file
    would otherwise break every parser at once.
    """
    if content.startswith(_OLE2_MAGIC):
        return "xls"
    if content.startswith(_ZIP_MAGIC):
        return "xlsx"
    return suffix.lower().lstrip(".")


def load(content: bytes, suffix: str = "") -> dict[str, Grid]:
    """Every sheet of a workbook, keyed by sheet name, in file order."""
    detected = sniff(content, suffix)
    if detected == "xls":
        return _load_xls(content)
    if detected in {"xlsx", "xlsm"}:
        return _load_xlsx(content)
    raise WorkbookError(f"unsupported workbook format {detected or suffix!r}")


def _load_xls(content: bytes) -> dict[str, Grid]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    grids: dict[str, Grid] = {}
    for sheet in book.sheets():
        rows = tuple(
            tuple(sheet.row_values(index)) for index in range(sheet.nrows)
        )
        grids[sheet.name] = Grid(name=sheet.name, rows=rows)
    return grids


def _load_xlsx(content: bytes) -> dict[str, Grid]:
    import openpyxl

    book = openpyxl.load_workbook(
        io.BytesIO(content), read_only=True, data_only=True
    )
    grids: dict[str, Grid] = {}
    try:
        for name in book.sheetnames:
            sheet = book[name]
            # Read-only mode trusts the sheet's declared ``dimension``, and
            # exported workbooks do get that wrong -- the PIHPS files declare
            # ``A1:C1`` for a 32x39 grid.  Resetting it makes the iterator read
            # what is actually there, so a bad declaration cannot silently
            # truncate a sheet into a plausible-looking empty parse.
            if hasattr(sheet, "reset_dimensions"):
                sheet.reset_dimensions()
            rows = tuple(tuple(row) for row in sheet.iter_rows(values_only=True))
            grids[name] = Grid(name=name, rows=rows)
    finally:
        book.close()
    return grids


@dataclass(frozen=True, slots=True)
class Period:
    """One data column: which date it is, and at what frequency."""

    col: int
    ref_date: dt.date
    ref_period: str
    label: str


def _period_number(grid: Grid, period_row: int, col: int) -> int | None:
    """Month or quarter number in a header cell, whichever it carries."""
    label = grid.label(period_row, col)
    month = parse_month(label)
    return month if month is not None else parse_quarter(label)


def _runs(grid: Grid, period_row: int, first_col: int) -> list[list[int]]:
    """Period-labelled columns, split into one block per year.

    A block ends where the period sequence stops rising — Des followed by Jan,
    Q4 by Q1 — and not merely where a blank column falls.  Both rules are
    needed: SEKI separates its quarter blocks with the year's own total column,
    while the consumer survey runs Jan-Des Jan-Des across 176 columns with
    nothing in between, and splitting only on blanks would read the whole sheet
    as one year.
    """
    runs: list[list[int]] = []
    current: list[int] = []
    previous: int | None = None
    for col in range(first_col, grid.ncols):
        number = _period_number(grid, period_row, col)
        if number is None:
            if current:
                runs.append(current)
                current, previous = [], None
            continue
        if previous is not None and number <= previous:
            runs.append(current)
            current = []
        current.append(col)
        previous = number
    if current:
        runs.append(current)
    return runs


#: How far before a block its year header may sit.  BI writes the year over
#: the block's own annual-total column, which precedes the months.
_YEAR_LOOKBACK = 2


def periods(
    grid: Grid,
    year_row: int,
    period_row: int,
    *,
    first_col: int = 0,
) -> list[Period]:
    """Map data columns to dates from a year row and a period row.

    Dating is by **block**, not by forward-filling the year across columns,
    because the year header does not reliably sit above the block it labels.
    In SEKI table VII.3 the years 2010 to 2020 line up with their quarters and
    2021, 2024 and 2025 sit three columns to the right of theirs; forward-fill
    therefore hands three of 2021's quarters to 2020 and stores two values for
    every quarter of six different years.  The same class of misalignment is
    already handled explicitly for OJK's by-province tables.

    So: consecutive period labels are grouped into blocks, and each block takes
    the year written within it (or just before it, where the block's own
    annual-total column carries the year).  A block with no year of its own, or
    one whose year would not move forwards, continues the sequence from the
    block before it.
    """
    year_cells = {
        col: parse_year(grid.cell(year_row, col))
        for col in range(first_col, grid.ncols)
        if parse_year(grid.cell(year_row, col)) is not None
    }
    out: list[Period] = []
    year: int | None = None
    for run in _runs(grid, period_row, first_col):
        candidates = [
            found
            for col, found in year_cells.items()
            if run[0] - _YEAR_LOOKBACK <= col <= run[-1]
        ]
        chosen = next((c for c in candidates if year is None or c > year), None)
        if chosen is not None:
            year = chosen
        elif year is not None:
            year += 1
        else:
            continue  # a block before any year header: nothing to date it by
        for col in run:
            label = grid.label(period_row, col)
            month = parse_month(label)
            quarter = parse_quarter(label)
            if month is not None:
                out.append(Period(col, month_start(year, month), "M", label))
            elif quarter is not None:
                out.append(Period(col, quarter_start(year, quarter), "Q", label))
    return out


def detect_header(
    grid: Grid, *, search_rows: int = 12, min_periods: int = 8
) -> tuple[int, int]:
    """Find the (year row, period row) pair without hard-coding positions."""
    best: tuple[int, tuple[int, int]] | None = None
    for year_row in range(min(search_rows, grid.nrows)):
        for period_row in range(year_row + 1, min(year_row + 4, grid.nrows)):
            found = len(periods(grid, year_row, period_row))
            if found >= min_periods and (best is None or found > best[0]):
                best = (found, (year_row, period_row))
    if best is None:
        raise WorkbookError(
            f"sheet {grid.name!r}: no year/period header rows in the first "
            f"{search_rows} rows"
        )
    return best[1]


def number(value: Any) -> float | None:
    """A numeric cell, or ``None`` for the dashes and blanks these tables use."""
    if value is None or isinstance(value, bool):
        return None
    if _is_number(value):
        return None if value != value else float(value)  # NaN guard
    text = str(value).strip().replace(" ", "")
    if not text or text in {"-", "–", "—", "...", "n.a.", "na", "#DIV/0!", "#N/A"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()").replace(" ", "")
    # Whichever of "." and "," comes last is the decimal separator: BI writes
    # 1,234.56 and OJK sometimes writes 1.234,56 in the same workbook family.
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        groups = text.split(",")
        thousands = all(len(part) == 3 for part in groups[1:])
        text = text.replace(",", "") if thousands else text.replace(",", ".")
    try:
        result = float(text)
    except ValueError:
        return None
    return -result if negative else result


def data_rows(grid: Grid, start: int, *, label_cols: int = 6) -> Iterator[tuple[int, str]]:
    """Rows from ``start`` that carry a component label, with their index."""
    for row in range(start, grid.nrows):
        label = grid.row_label(row, label_cols)
        if label:
            yield row, label


def month_year_periods(
    grid: Grid,
    label_row: int,
    *,
    sub_row: int | None = None,
    sub_label: str = "",
    first_col: int = 0,
) -> list[Period]:
    """Columns headed by a full period label such as ``'Desember 2024'``.

    OJK's snapshot tables head each block with a month-and-year string rather
    than the year/month pair BI uses, and split the block into a nominal column
    and a share column.  ``sub_label`` picks which of those to read; the label
    row is forward-filled across the block, the same way a merged cell reads.
    """
    out: list[Period] = []
    current: tuple[dt.date, str] | None = None
    for col in range(first_col, grid.ncols):
        label = grid.label(label_row, col)
        if label:
            month = None
            year = parse_year(label)
            parts = label.replace(",", " ").split()
            for part in parts:
                month = parse_month(part) or month
            if year is not None and month is not None:
                current = (month_start(year, month), label)
            elif year is not None:
                current = (dt.date(year, 1, 1), label)
            else:
                current = None
        if current is None:
            continue
        if sub_row is not None:
            sub = grid.label(sub_row, col)
            if sub_label and sub_label.casefold() not in sub.casefold():
                continue
        ref_date, text = current
        period = "M" if ref_date.month != 1 or " " in text.strip() else "A"
        out.append(Period(col, ref_date, "M" if period == "M" else "A", text))
    return out


def rolling_month_columns(
    grid: Grid,
    month_row: int,
    latest: dt.date,
    *,
    first_col: int = 0,
) -> list[Period]:
    """Date the month columns backwards from a known latest month.

    OJK's by-province tables carry a rolling window of months whose merged year
    headers do not line up with the month sequence — in the June 2025 release
    the "2024" merge runs three months into 2025.  The release month is known
    from the page that published the workbook, so walking right to left from it
    is both simpler and correct, where trusting the merge is neither.
    """
    columns = [
        col
        for col in range(first_col, grid.ncols)
        if parse_month(grid.label(month_row, col)) is not None
    ]
    out: list[Period] = []
    year, month = latest.year, latest.month
    for col in reversed(columns):
        out.append(Period(col, month_start(year, month), "M", grid.label(month_row, col)))
        month -= 1
        if month == 0:
            year, month = year - 1, 12
    out.reverse()
    return out


def annual_columns(
    grid: Grid, year_row: int, month_row: int, *, first_col: int = 0
) -> list[Period]:
    """Columns headed by a year alone — a year-end position, not a month."""
    out: list[Period] = []
    for col in range(first_col, grid.ncols):
        year = parse_year(grid.cell(year_row, col))
        if year is None:
            continue
        if parse_month(grid.label(month_row, col)) is not None:
            continue
        out.append(Period(col, dt.date(year, 1, 1), "A", str(year)))
    return out


def find_month_row(grid: Grid, *, limit: int = 8, minimum: int = 4) -> int | None:
    """The header row carrying month labels, discovered rather than assumed.

    OJK's DPK sheets put it at row index 2 today; a title line added above
    would move it, and a fixed index fails silently by reading the wrong row.
    """
    best: tuple[int, int] | None = None
    for row in range(min(limit, grid.nrows)):
        found = sum(
            1
            for col in range(grid.ncols)
            # Only a written month name counts.  parse_month also accepts a
            # bare 1-12, which a row of small data values would satisfy — and
            # a data row mistaken for the header dates the whole sheet wrong.
            if isinstance(grid.cell(row, col), str)
            and parse_month(grid.label(row, col)) is not None
        )
        if found >= minimum and (best is None or found > best[0]):
            best = (found, row)
    return None if best is None else best[1]


def find_label_row(grid: Grid, *needles: str, limit: int = 8) -> int | None:
    """The header row that mentions every one of ``needles``."""
    wanted = [needle.casefold() for needle in needles]
    for row in range(min(limit, grid.nrows)):
        joined = " ".join(grid.label(row, col) for col in range(grid.ncols)).casefold()
        if all(needle in joined for needle in wanted):
            return row
    return None
