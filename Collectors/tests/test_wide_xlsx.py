"""The wide-workbook sink: merge weekly columns into the PIHPS files."""

from __future__ import annotations

import datetime as dt
import shutil

import pytest

from collectors.paths import FOOD_PRICES
from collectors.sinks import wide_xlsx
from collectors.sinks.wide_xlsx import Grid, Row
from collectors.sources import pihps

REAL = sorted(FOOD_PRICES.rglob("*.xlsx"))
LIVE_2026 = FOOD_PRICES / "Traditional Market" / "Tabel Harga Berdasarkan Daerah 2026.xlsx"


def test_header_date_uses_the_on_disk_spelling():
    """The API sends 01/01/2026; the workbooks spell it 01/ 01/ 2026."""
    assert wide_xlsx.header_date(dt.date(2026, 1, 1)) == "01/ 01/ 2026"
    assert wide_xlsx.header_date(dt.date(2026, 12, 31)) == "31/ 12/ 2026"


@pytest.mark.parametrize("text", ["01/ 01/ 2026", "01/01/2026", " 1/ 1/ 2026 "])
def test_both_date_spellings_parse(text):
    assert wide_xlsx.parse_header_date(text) == dt.date(2026, 1, 1)


@pytest.mark.parametrize("text,expected", [("16,200", True), ("-", False), ("", False), (None, False)])
def test_is_value_treats_a_dash_as_missing(text, expected):
    assert wide_xlsx.is_value(text) is expected


@pytest.mark.parametrize("source", REAL, ids=lambda p: f"{p.parent.name}-{p.name.split()[-1]}")
def test_every_real_workbook_reads_31_commodities(source):
    grid = wide_xlsx.read_grid(source)
    assert len(grid.rows) == 31
    assert grid.weeks, "no date columns found"
    # A stale `dimension` in these files makes openpyxl's read-only mode yield
    # almost nothing, which would blank the workbook on merge.
    assert len(grid.values) >= 31 * len(grid.weeks) * 0.9


@pytest.mark.parametrize("source", REAL, ids=lambda p: f"{p.parent.name}-{p.name.split()[-1]}")
def test_read_write_read_is_lossless(source, tmp_path):
    grid = wide_xlsx.read_grid(source)
    target = tmp_path / source.name
    wide_xlsx.write_grid(target, grid)
    assert wide_xlsx.read_grid(target).snapshot() == grid.snapshot()


def test_the_hierarchy_and_exact_labels_survive(tmp_path):
    grid = wide_xlsx.read_grid(LIVE_2026)
    target = tmp_path / "out.xlsx"
    wide_xlsx.write_grid(target, grid)
    rows = wide_xlsx.read_grid(target).rows
    assert [r.no for r in rows[:7]] == ["I", "1", "2", "3", "4", "5", "6"]
    assert rows[0].name == "Beras" and rows[0].level == 1
    assert rows[1].name == "Beras Kualitas Bawah I" and rows[1].level == 2
    # PIHPS publishes this label with a trailing space and so does the file.
    assert any(r.name == "Cabai Merah Keriting " for r in rows)


def test_the_sheet_is_named_Sheet_and_the_header_row_is_frozen(tmp_path):
    import openpyxl

    target = tmp_path / "out.xlsx"
    wide_xlsx.write_grid(target, wide_xlsx.read_grid(LIVE_2026))
    book = openpyxl.load_workbook(target)
    assert book.sheetnames == ["Sheet"]
    assert book["Sheet"].freeze_panes == "A2"


def test_prices_are_written_as_text_with_thousands_commas(tmp_path):
    import openpyxl

    target = tmp_path / "out.xlsx"
    wide_xlsx.write_grid(target, wide_xlsx.read_grid(LIVE_2026))
    sheet = openpyxl.load_workbook(target)["Sheet"]
    value = sheet.cell(row=2, column=3).value
    assert isinstance(value, str) and "," in value


def _grid(week: dt.date, text: str, name: str = "Beras") -> Grid:
    row = Row(no="I", name=name, level=1)
    return Grid(rows=[row], values={(row.key, week): text})


def test_a_new_week_is_added_and_reported():
    existing = _grid(dt.date(2026, 9, 3), "16,000")
    fetched = _grid(dt.date(2026, 9, 10), "16,100")
    merged, report = wide_xlsx.merge(existing, fetched, path="x.xlsx")
    assert report.added_weeks == [dt.date(2026, 9, 10)]
    assert merged.weeks == [dt.date(2026, 9, 3), dt.date(2026, 9, 10)]
    assert merged.cell(merged.rows[0], dt.date(2026, 9, 3)) == "16,000"


def test_a_dash_never_clobbers_an_existing_price():
    """A dash means the market did not report, not that the price was wrong."""
    existing = _grid(dt.date(2026, 9, 3), "16,000")
    fetched = _grid(dt.date(2026, 9, 3), "-")
    merged, report = wide_xlsx.merge(existing, fetched, path="x.xlsx")
    assert merged.cell(merged.rows[0], dt.date(2026, 9, 3)) == "16,000"
    assert report.ignored_dashes == 1
    assert not report.changed


def test_a_dash_is_kept_where_there_is_nothing_to_protect():
    existing = Grid()
    fetched = _grid(dt.date(2026, 9, 3), "-")
    merged, _ = wide_xlsx.merge(existing, fetched, path="x.xlsx")
    assert merged.cell(merged.rows[0], dt.date(2026, 9, 3)) == "-"


def test_a_real_value_fills_a_gap_left_by_an_earlier_dash():
    existing = _grid(dt.date(2026, 9, 3), "-")
    fetched = _grid(dt.date(2026, 9, 3), "16,000")
    merged, report = wide_xlsx.merge(existing, fetched, path="x.xlsx")
    assert merged.cell(merged.rows[0], dt.date(2026, 9, 3)) == "16,000"
    assert report.filled == 1


def test_a_revised_price_is_applied_and_reported():
    existing = _grid(dt.date(2026, 9, 3), "16,000")
    fetched = _grid(dt.date(2026, 9, 3), "16,050")
    merged, report = wide_xlsx.merge(existing, fetched, path="x.xlsx")
    assert merged.cell(merged.rows[0], dt.date(2026, 9, 3)) == "16,050"
    assert report.revised == [("Beras", dt.date(2026, 9, 3), "16,000", "16,050")]


def test_a_week_the_api_no_longer_returns_is_retained():
    existing = _grid(dt.date(2026, 1, 1), "15,000")
    fetched = _grid(dt.date(2026, 9, 3), "16,000")
    merged, _ = wide_xlsx.merge(existing, fetched, path="x.xlsx")
    assert dt.date(2026, 1, 1) in merged.weeks


def test_rows_match_by_name_not_position():
    """A reordered payload must still land on the right commodity."""
    a, b = Row("I", "Beras", 1), Row("II", "Daging Ayam", 1)
    week = dt.date(2026, 9, 3)
    existing = Grid(rows=[a, b], values={(a.key, week): "16,000", (b.key, week): "39,000"})
    fetched = Grid(rows=[b, a], values={(a.key, week): "16,100", (b.key, week): "39,100"})
    merged, _ = wide_xlsx.merge(existing, fetched, path="x.xlsx")
    assert [r.name for r in merged.rows] == ["Beras", "Daging Ayam"]
    assert merged.cell(a, week) == "16,100" and merged.cell(b, week) == "39,100"


def test_an_unknown_commodity_is_appended_and_reported():
    existing = _grid(dt.date(2026, 9, 3), "16,000")
    fetched = _grid(dt.date(2026, 9, 3), "1,000", name="Kedelai")
    merged, report = wide_xlsx.merge(existing, fetched, path="x.xlsx")
    assert [r.name for r in merged.rows] == ["Beras", "Kedelai"]
    assert report.added_rows == ["Kedelai"]


def test_an_unchanged_merge_does_not_rewrite_the_file(tmp_path):
    """openpyxl output is not reproducible byte-wise, so the guard is the grid."""
    target = tmp_path / "out.xlsx"
    shutil.copy2(LIVE_2026, target)
    before = target.read_bytes()
    report = wide_xlsx.write(target, wide_xlsx.read_grid(LIVE_2026))
    assert not report.written and not report.changed
    assert target.read_bytes() == before


def test_a_second_run_after_a_real_change_is_a_no_op(tmp_path):
    target = tmp_path / "out.xlsx"
    shutil.copy2(LIVE_2026, target)
    fetched = _grid(dt.date(2026, 9, 17), "16,400")
    assert wide_xlsx.write(target, fetched).written
    after_first = target.read_bytes()
    assert not wide_xlsx.write(target, fetched).written
    assert target.read_bytes() == after_first


def test_dry_run_reports_without_writing(tmp_path):
    target = tmp_path / "out.xlsx"
    shutil.copy2(LIVE_2026, target)
    before = target.read_bytes()
    report = wide_xlsx.write(target, _grid(dt.date(2026, 9, 17), "16,400"), dry_run=True)
    assert report.added_weeks == [dt.date(2026, 9, 17)] and not report.written
    assert target.read_bytes() == before


def test_merging_a_recorded_payload_into_the_real_workbook_changes_nothing(pihps_grid, tmp_path):
    """The uploaded files are a capture of this endpoint, so a replay is a no-op."""
    target = tmp_path / "out.xlsx"
    shutil.copy2(LIVE_2026, target)
    fetched = pihps.grid_from_response(pihps_grid)
    report = wide_xlsx.write(target, fetched)
    assert not report.changed, report.describe()


def test_the_recorded_payload_parses_into_the_expected_shape(pihps_grid):
    grid = pihps.grid_from_response(pihps_grid)
    assert len(grid.rows) == 7
    assert len(grid.weeks) == 6
    assert grid.weeks[-1] == dt.date(2026, 9, 10)
    assert any(r.name == "Cabai Merah Keriting " for r in grid.rows)
    assert {r.level for r in grid.rows} == {1, 2}
